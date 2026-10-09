from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import Optional
import io

from database import get_db
from auth.dependencies import require_management_or_above
from services import placement_rate, salary_analytics, branch_analytics

router = APIRouter(prefix="/export", tags=["Export"])


def _get_module_data(module: str, db: Session, year: Optional[int]):
    if module == "placement-rate":
        return placement_rate.get_placement_rate(db, year)
    elif module == "salary":
        return salary_analytics.get_salary_analytics(db, year)
    elif module == "branch":
        return branch_analytics.get_branch_analytics(db, year)
    return {}


@router.get("/excel")
def export_excel(
    module: str = Query(..., description="placement-rate | salary | branch"),
    year: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _=Depends(require_management_or_above),
):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    data = _get_module_data(module, db, year)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = module.replace("-", " ").title()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def write_section(title, rows, headers):
        # Section title
        ws.append([title])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=13)
        ws.append([])
        # Headers
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        ws.append([])

    if module == "placement-rate":
        write_section(
            "Overall Stats",
            [[data["overall_rate"], data["total_students"], data["placed_students"], data["unplaced_students"]]],
            ["Placement Rate %", "Total", "Placed", "Unplaced"],
        )
        write_section(
            "Branch Breakdown",
            [[b["branch"], b["total"], b["placed"], b["rate"]] for b in data.get("branch_breakdown", [])],
            ["Branch", "Total", "Placed", "Rate %"],
        )

    elif module == "salary":
        overall = data.get("overall", {})
        write_section(
            "Overall Salary Stats",
            [[overall.get("min"), overall.get("max"), overall.get("avg"), overall.get("median"), overall.get("total_offers")]],
            ["Min LPA", "Max LPA", "Avg LPA", "Median LPA", "Total Offers"],
        )
        write_section(
            "Branch Salary",
            [[b["branch"], b["avg"], b["min"], b["max"], b["count"]] for b in data.get("branch_salary", [])],
            ["Branch", "Avg LPA", "Min LPA", "Max LPA", "Count"],
        )
        write_section(
            "Top Companies",
            [[c["company"], c["avg_package"]] for c in data.get("top_companies", [])],
            ["Company", "Avg Package LPA"],
        )

    elif module == "branch":
        write_section(
            "Branch Analytics",
            [[b["branch"], b["total_students"], b["placed"], b["placement_rate"], b["avg_package"], b["companies_visited"]]
             for b in data.get("branch_data", [])],
            ["Branch", "Total", "Placed", "Rate %", "Avg Package", "Companies Visited"],
        )

    # Auto-width
    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"placement_{module}_{year or 'all'}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/pdf")
def export_pdf(
    module: str = Query(..., description="placement-rate | salary | branch"),
    year: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _=Depends(require_management_or_above),
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    data = _get_module_data(module, db, year)
    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    title = f"College Placement Analytics — {module.replace('-', ' ').title()}"
    if year:
        title += f" ({year})"
    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Spacer(1, 12))

    def make_table(header_row, body_rows):
        table_data = [header_row] + body_rows
        t = Table(table_data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EBF3FF")]),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        return t

    if module == "placement-rate":
        elements.append(Paragraph("Overall Stats", styles["Heading2"]))
        elements.append(make_table(
            ["Placement Rate %", "Total Students", "Placed", "Unplaced"],
            [[data["overall_rate"], data["total_students"], data["placed_students"], data["unplaced_students"]]]
        ))
        elements.append(Spacer(1, 16))
        elements.append(Paragraph("Branch Breakdown", styles["Heading2"]))
        elements.append(make_table(
            ["Branch", "Total", "Placed", "Rate %"],
            [[b["branch"], b["total"], b["placed"], b["rate"]] for b in data.get("branch_breakdown", [])]
        ))

    elif module == "salary":
        overall = data.get("overall", {})
        elements.append(Paragraph("Overall Salary Stats", styles["Heading2"]))
        elements.append(make_table(
            ["Min LPA", "Max LPA", "Avg LPA", "Median LPA", "Total Offers"],
            [[overall.get("min"), overall.get("max"), overall.get("avg"), overall.get("median"), overall.get("total_offers")]]
        ))
        elements.append(Spacer(1, 16))
        elements.append(Paragraph("Branch Salary", styles["Heading2"]))
        elements.append(make_table(
            ["Branch", "Avg LPA", "Min LPA", "Max LPA", "Count"],
            [[b["branch"], b["avg"], b["min"], b["max"], b["count"]] for b in data.get("branch_salary", [])]
        ))

    elif module == "branch":
        elements.append(Paragraph("Branch Analytics", styles["Heading2"]))
        elements.append(make_table(
            ["Branch", "Total", "Placed", "Rate %", "Avg Pkg", "Companies"],
            [[b["branch"], b["total_students"], b["placed"], b["placement_rate"], b["avg_package"], b["companies_visited"]]
             for b in data.get("branch_data", [])]
        ))

    doc.build(elements)
    stream.seek(0)
    filename = f"placement_{module}_{year or 'all'}.pdf"
    return StreamingResponse(
        stream,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
